"""조직도 Excel을 로컬 서버 전용 JSON으로 변환합니다. 원본 값은 로그에 출력하지 않습니다."""
import sys, json, os, re
from pathlib import Path
from datetime import datetime, timezone
from openpyxl import load_workbook

def convert(source, target):
    source, target = Path(source), Path(target)
    if not source.is_file(): raise ValueError('조직도 원본 파일을 찾지 못했습니다.')
    if source.suffix.lower() != '.xlsx': raise ValueError('.xlsx 파일을 사용해 주세요.')
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    required = ['번호','부서','이름','사무실전화번호','담당업무','팀명','사용자정렬순서']
    employees, found, skipped = [], False, 0
    try:
        for sheet in workbook:
            headers = None
            for row in sheet.iter_rows():
                values = [str(c.value or '').strip() for c in row]
                if headers is None:
                    normalized = [re.sub(r'\s+', '', v) for v in values]
                    if all(key in normalized for key in required):
                        headers = {key:normalized.index(key) for key in required}; found = True
                    continue
                if not any(values): continue
                if any(row[i].data_type == 'f' for i in headers.values() if i < len(row)):
                    raise ValueError('필수 항목에 수식이 있습니다. 값을 확인하여 값으로 저장한 Excel을 사용해 주세요.')
                get = lambda key: values[headers[key]] if headers[key] < len(values) else ''
                if not get('부서') or not get('담당업무'): skipped += 1; continue
                phone = get('사무실전화번호')
                cell = row[headers['사무실전화번호']]
                if isinstance(cell.value, (int,float)) and re.fullmatch(r'0+',cell.number_format or ''):
                    phone = str(int(cell.value)).zfill(len(cell.number_format))
                employees.append(dict(id=get('번호'),department=get('부서'),name=get('이름'),phone=phone,duties=get('담당업무'),team=get('팀명'),sortOrder=get('사용자정렬순서')))
    finally: workbook.close()
    if not found: raise ValueError('필수 컬럼 7개를 찾지 못했습니다. 헤더명을 확인해 주세요.')
    if not employees: raise ValueError('사용 가능한 담당업무 행이 없습니다. 기존 JSON은 유지합니다.')
    payload = dict(schemaVersion=1,updatedAt=datetime.now(timezone.utc).isoformat(),employees=employees)
    target.parent.mkdir(parents=True,exist_ok=True)
    temporary = target.with_suffix(target.suffix+'.tmp')
    temporary.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    os.replace(temporary,target)
    print(json.dumps(dict(converted=len(employees),skipped=skipped),ensure_ascii=True))
if __name__ == '__main__':
    try:
        if len(sys.argv) != 3: raise ValueError('입력 xlsx 경로와 출력 JSON 경로가 필요합니다.')
        convert(sys.argv[1],sys.argv[2])
    except Exception as error:
        # 내부 셀/담당자 값을 포함할 수 있는 예외 원문은 출력하지 않습니다.
        if isinstance(error,ValueError): print(str(error),file=sys.stderr)
        else: print('조직도 변환에 실패했습니다. 파일 형식·접근 권한·Python openpyxl 설치를 확인하세요.',file=sys.stderr)
        sys.exit(1)
